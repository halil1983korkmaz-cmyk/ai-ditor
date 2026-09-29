"""Separate summaries and issue numbering for the social-science preset."""
import copy
import io
import unittest
from docx import Document
from docx.oxml.ns import qn
from docx_export import generate_docx_from_form
from formatter import generate_latex_from_form
from journal_templates import TEMPLATES, normalize_settings
from page_furniture import article_values, running_slots
from test_formatter_templates import sample_article


def settings(**overrides):
    return normalize_settings({**next(t['settings'] for t in TEMPLATES if t['id']=='scholarly'),
                               'journal_name_tr':'Örnek Sosyal Bilimler',
                               'journal_name_en':'Example Social Sciences', **overrides})


class ScholarlyTemplateTests(unittest.TestCase):
    def test_separate_summary_has_continuous_numbering_and_native_headers(self):
        data=sample_article();data['cover']['start_page']='88'
        doc=Document(io.BytesIO(generate_docx_from_form(data,{},settings())))
        self.assertEqual(len(doc.sections),2)
        cover,body=doc.sections
        self.assertAlmostEqual(cover.left_margin.cm,2.5,places=2)
        self.assertTrue(cover.different_first_page_header_footer)
        self.assertFalse(body.different_first_page_header_footer)
        self.assertEqual(cover._sectPr.find(qn('w:pgNumType')).get(qn('w:start')),'88')
        self.assertIsNone(body._sectPr.find(qn('w:pgNumType')))
        self.assertIn('Örnek & Smith',body.header.tables[0].cell(0,0).text)
        self.assertIn('88–104',body.even_page_header.tables[0].cell(0,0).text)
        self.assertIn('Gönderim / Received', '\n'.join(p.text for p in cover.first_page_footer.paragraphs))
        text=[p.text for p in doc.paragraphs]
        self.assertLess(text.index(data['abstract']['tr_abs']),text.index(data['cover']['en_title']))
        self.assertLess(text.index('Extended Summary'),text.index(data['abstract']['en_abs']))
        self.assertLess(text.index(data['abstract']['en_abs']),text.index(data['sections'][0]['name']))
        self.assertEqual(text.count(data['cover']['en_title']),1)
        self.assertTrue(body.header.is_linked_to_previous)

    def test_long_summary_is_editable_and_can_flow_across_pages(self):
        data=sample_article();paragraphs=[f'Summary paragraph {i}: '+data['abstract']['en_abs'] for i in range(18)]
        data['abstract']['en_abs']='\n\n'.join(paragraphs)
        doc=Document(io.BytesIO(generate_docx_from_form(data,{},settings())))
        found=[p for p in doc.paragraphs if p.text.startswith('Summary paragraph')]
        self.assertEqual([p.text for p in found],paragraphs)
        self.assertTrue(all(p.paragraph_format.keep_with_next is not True for p in found))
        self.assertTrue(all(p.paragraph_format.line_spacing==1.5 for p in found))
        self.assertEqual(len(doc.sections),2)

    def test_english_only_does_not_duplicate_summary_or_turkish_labels(self):
        data=sample_article()
        doc=Document(io.BytesIO(generate_docx_from_form(data,{},settings(english_only=True))))
        text='\n'.join(p.text for p in doc.paragraphs)
        self.assertEqual(text.count(data['abstract']['en_abs']),1)
        self.assertNotIn(data['abstract']['tr_abs'],text)
        self.assertNotIn('\nÖz\n',text)
        self.assertIn('Extended Summary',text)
        tex=generate_latex_from_form(data,{},settings(english_only=True))
        layout=tex[tex.index('% Template:'):]
        self.assertNotIn('\\JGTTRturkishabstract',layout)
        self.assertEqual(layout.count('\\JGTTRenglishabstract'),1)

    def test_tex_layout_splits_summary_and_preserves_customization(self):
        data=sample_article();prefs=settings(english_abstract_heading=r'Overview & findings',footer_first_mode='custom',footer_first_center='CUSTOM {sayfa}',header_left='CUSTOM {yazarlar}')
        tex=generate_latex_from_form(data,{},prefs)
        layout=tex[tex.index('% Template:'):tex.index('\\begin{document}')]
        self.assertLess(layout.index('\\restoregeometry'),layout.index('\\JGTTRenglishabstract'))
        self.assertIn('Overview \\& findings',layout)
        self.assertIn('left=2.5cm,right=2.5cm',tex)
        self.assertIn(r'\setlength{\parindent}{1.25cm}',tex)
        self.assertIn(r'\renewcommand{\baselinestretch}{1.5}',tex)
        self.assertIn(r'CUSTOM \thepage{}',tex)
        self.assertIn('CUSTOM Örnek',tex)
        self.assertEqual(article_values(data,prefs)['sayfa_araligi'],'89–104')
        self.assertEqual(running_slots(prefs,data,'header','odd')[2],'{sayfa}')

    def test_single_language_cover_still_transitions_to_body_without_restart(self):
        data=sample_article();data['abstract']['en_abs']=''
        doc=Document(io.BytesIO(generate_docx_from_form(data,{},settings(doi_position='bottom'))))
        self.assertIn(data['cover']['doi'],doc.sections[0].first_page_footer._element.xml)
        self.assertNotIn(data['cover']['doi'],doc.element.xml)
        self.assertEqual(len(doc.sections),2)
        self.assertIsNone(doc.sections[1]._sectPr.find(qn('w:pgNumType')))
        self.assertNotIn('Extended Summary',[p.text for p in doc.paragraphs])
        table=doc.tables[-1]
        self.assertAlmostEqual(sum(c.width.cm for c in table.columns),16,places=2)
        body=next(p for p in doc.paragraphs if p.text.startswith('Editoryal çalışmalar'))
        self.assertAlmostEqual(body.paragraph_format.first_line_indent.cm,1.25,places=2)
