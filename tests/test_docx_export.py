"""Editable output, page variants, safe citation targets and authenticated delivery."""
import base64
import copy
import io
import json
from pathlib import Path
import tempfile
import unittest
import zipfile
from types import SimpleNamespace
from unittest.mock import Mock, patch

from docx import Document
from lxml import etree
from pypdf import PdfWriter
from docx_export import generate_docx_from_form, DOCX_MIME
from formatter import generate_latex_from_form
from journal_templates import TEMPLATES, normalize_settings
from citation_links import CitationIndex, citation_report
from page_furniture import running_slots
from test_formatter_templates import sample_article
from test_accounts import make_setup, register, headers, PNG
from desktop import ArticleDownloads

NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

class CitationTests(unittest.TestCase):
    def test_parenthetical_narrative_multiple_and_suffixes(self):
        refs=['Yılmaz, A. (2020a). First.', 'Yılmaz, A. (2020b). Second.', 'Kaya, B., & Demir, C. (2021). Third.', 'Smith, A., Jones, B., & Brown, C. (2022). Fourth.', 'World Health Organization. (2024). Fifth.']
        text='Yılmaz (2020a); (Yılmaz, 2020b; Kaya ve Demir, 2021, s. 4); Smith et al. (2022); World Health Organization (2024).'
        spans=list(CitationIndex(refs).spans(text))
        self.assertEqual([i for _,_,i in spans],[0,1,2,3,4])
        self.assertEqual(text[spans[0][0]:spans[0][1]],'Yılmaz (2020a)')
        self.assertEqual(list(CitationIndex(refs).spans('Yılmaz (2020), 2024 ve Smith (2022)')),[])

    def test_ambiguity_and_single_author_tail_never_linked(self):
        index=CitationIndex(['Kaya, B. (2020). A.', 'Kaya, C. (2020). B.', 'Demir, D. (2021). C.'])
        self.assertEqual(list(index.spans('Kaya (2020); Kaya ve Demir (2021); (Kaya & Demir, 2021)')),[])
        self.assertEqual(len(index.ambiguous),2)
        self.assertEqual(list(CitationIndex(['Örnek, A. (2020). A.'],False).spans('Örnek (2020)')),[])

    def test_exact_bookmark_targets_and_original_text(self):
        data=sample_article();data['references']='Örnek, D. (2026). Örnek. https://doi.org/10.1234/example\nAdams, A. (2020). Other.'
        text='Örnek (2026), içerik & veri güvenliğini tartışmıştır. (Adams, 2020).'
        data['sections'][0]['content']=text
        blob=generate_docx_from_form(data,{})
        with zipfile.ZipFile(io.BytesIO(blob)) as z:
            root=etree.fromstring(z.read('word/document.xml'))
            links=root.xpath('//w:hyperlink[@w:anchor]',namespaces=NS)
            self.assertEqual([x.get('{'+NS['w']+'}anchor') for x in links],['aiditor_ref_1','aiditor_ref_0'])
            self.assertEqual(''.join(links[0].xpath('.//w:t/text()',namespaces=NS)), 'Örnek (2026)')
            self.assertEqual(root.xpath('//w:bookmarkStart/@w:name',namespaces=NS),['aiditor_ref_0','aiditor_ref_1'])
            self.assertIn('https://doi.org/10.1234/example',z.read('word/_rels/document.xml.rels').decode())
        tex=generate_latex_from_form(data,{})
        self.assertIn(r'\hyperlink{aiditor_ref_1}{Örnek (2026)}',tex)
        self.assertIn(r'\hypertarget{aiditor_ref_1}{}Örnek',tex)
        self.assertIn(r'\& veri',tex)
        self.assertEqual(citation_report(data)['linked_in_body'],2)

class WordOutputTests(unittest.TestCase):
    def test_cover_notes_are_bottom_anchored_and_do_not_repeat_on_body_pages(self):
        for template in TEMPLATES:
            with self.subTest(layout=template['id']):
                data=sample_article()
                settings={**template['settings'], 'footer_text':'Derginin sabit dipnotu',
                          'footer_first_mode':'custom','footer_first_center':'Kapak {sayfa}'}
                doc=Document(io.BytesIO(generate_docx_from_form(data,{},settings)))
                main='\n'.join(p.text for p in doc.paragraphs)
                footer=doc.sections[0].first_page_footer
                notes='\n'.join(p.text for p in footer.paragraphs)
                for text in ['Derginin sabit dipnotu',data['cover']['ethics'],data['cover']['title_note']]:
                    self.assertIn(text,notes)
                    self.assertNotIn(text,main)
                self.assertIn('Kapak',footer.tables[0].cell(0,0).text)
                self.assertEqual(len(doc.sections),2)
                self.assertFalse(doc.sections[1].different_first_page_header_footer)
                self.assertNotIn('Derginin sabit dipnotu',doc.sections[1].footer._element.xml)
                self.assertNotIn('w:start=',doc.sections[1]._sectPr.xml)

    def test_editable_templates_preserve_content(self):
        for template in TEMPLATES:
            with self.subTest(layout=template['id']):
                data=sample_article();blob=generate_docx_from_form(data,{},template['settings'],{'logo':('logo.png',PNG)})
                doc=Document(io.BytesIO(blob))
                self.assertEqual(doc.core_properties.title,data['cover']['tr_title'])
                self.assertTrue(any(p.text==data['cover']['tr_title']+' *' for p in doc.paragraphs))
                self.assertTrue(any('Aşama / Stage' in cell.text for t in doc.tables for row in t.rows for cell in row.cells))
                self.assertEqual(len(doc.inline_shapes),1)
                with zipfile.ZipFile(io.BytesIO(blob)) as z:self.assertIsNone(z.testzip())

    def test_running_variants_are_real_header_parts_and_page_fields(self):
        data=sample_article()
        settings=normalize_settings({'header_mode':'odd_even','header_left':'ODD {dergi}','header_even_left':'EVEN {yazarlar}','header_first_mode':'custom','header_first_center':'FIRST','footer_mode':'same','footer_center':'Page {sayfa}'})
        blob=generate_docx_from_form(data,{},settings)
        doc=Document(io.BytesIO(blob));section=doc.sections[0]
        self.assertTrue(doc.settings.odd_and_even_pages_header_footer)
        self.assertTrue(section.different_first_page_header_footer)
        self.assertIn('ODD',section.header.tables[0].cell(0,0).text)
        self.assertIn('EVEN',section.even_page_header.tables[0].cell(0,0).text)
        self.assertEqual(section.first_page_header.tables[0].cell(0,0).text,'FIRST')
        self.assertIn(' PAGE ',section.footer._element.xml)
        self.assertIn('w:start="89"',section._sectPr.xml)
        settings['header_first_mode']='inherit';data['cover']['start_page']='88'
        self.assertTrue(running_slots(settings,data,'header','first')[0].startswith('EVEN'))
        tex=generate_latex_from_form(data,{},settings)
        self.assertIn('a4paper,twoside',tex)
        self.assertIn(r'\fancyhead[CE]',tex)
        self.assertIn(r'\thepage{}',tex)
        self.assertNotIn('includehead=false',tex)

    def test_rich_merged_cells_keep_text_and_header_repeat(self):
        data=sample_article();data['figtables'][0]['tbl_model']={'rows':[[{'text':'Wide','colspan':2,'bgcolor':'#EEDDCC'}],[{'text':'Tall','rowspan':2},{'text':'A'}],[{'text':'B'}]],'header_rows':1,'column_widths':[.4,.6]}
        doc=Document(io.BytesIO(generate_docx_from_form(data,{})));table=doc.tables[-1]
        self.assertEqual(table.cell(0,0).text,'Wide');self.assertEqual(table.cell(0,1).text,'Wide')
        self.assertEqual(table.cell(1,0).text,'Tall');self.assertEqual(table.cell(2,0).text,'Tall')
        self.assertEqual(table.cell(2,1).text,'B');self.assertIn('w:tblHeader',table._tbl.xml)

    def test_pdf_assets_and_figure_are_rasterized(self):
        pdf=PdfWriter();pdf.add_blank_page(width=80,height=50);stream=io.BytesIO();pdf.write(stream)
        data=sample_article();data['figtables']=[{'type':'figure','file_key':'a','number':'1','tr_cap':'Görsel','section_id':'introduction'}]
        blob=generate_docx_from_form(data,{'a':('a.pdf',stream.getvalue())},assets={'logo':('logo.pdf',stream.getvalue())})
        self.assertEqual(len(Document(io.BytesIO(blob)).inline_shapes),2)

    def test_invalid_preferences_rejected_and_tex_escaped(self):
        for settings in ({'header_left':'{secret}'},{'header_left':'x'*181},{'footer_mode':'garbage'},{'link_citations':'yes'}):
            with self.assertRaises(ValueError):normalize_settings(settings)
        settings={'header_mode':'same','header_left':r'\input Evil & 10%'}
        tex=generate_latex_from_form(sample_article(),{},settings)
        self.assertNotIn(r'\input Evil',tex)
        self.assertIn(r'\& 10\%',tex)

class WordDeliveryTests(unittest.TestCase):
    def test_authenticated_isolated_api_and_native_bytes(self):
        with tempfile.TemporaryDirectory() as directory:
            app,_=make_setup(Path(directory));a,b=app.test_client(),app.test_client();ua,ub=register(a),register(b,'journal_b')
            settings={'header_mode':'odd_even','header_left':'TEK','header_even_left':'ÇİFT','header_first_mode':'custom','header_first_center':'KAPAK'}
            response=a.put('/api/journal',json={'settings':settings,'base_revision':0},headers=headers(ua));self.assertEqual(response.status_code,200)
            response=a.post('/process_docx',data={'data':json.dumps(sample_article())},headers=headers(ua));self.assertEqual(response.status_code,200,response.json)
            url='/download_docx/'+response.json['key'];download=a.get(url);self.assertEqual(download.mimetype,DOCX_MIME)
            doc=Document(io.BytesIO(download.data));self.assertIn('ÇİFT',doc.sections[0].even_page_header.tables[0].cell(0,0).text)
            self.assertEqual(b.get(url).status_code,404);self.assertEqual(app.test_client().get(url).status_code,401)
            self.assertEqual(a.get('/download_docx/expired').status_code,404)
            self.assertEqual(a.get('/api/journal',headers=headers(ua)).json['settings']['header_left'],'TEK')
            target=Path(directory)/'makale.docx';bridge=ArticleDownloads();bridge._window=SimpleNamespace(create_file_dialog=Mock(return_value=[str(target)]))
            webview=SimpleNamespace(FileDialog=SimpleNamespace(SAVE=30))
            with patch.dict('sys.modules',{'webview':webview}):result=bridge.save_article_docx(base64.b64encode(download.data).decode())
            self.assertEqual(result,{'ok':True,'cancelled':False});self.assertEqual(target.read_bytes(),download.data)
            self.assertFalse(bridge.save_article_docx(base64.b64encode(b'{"error":"login"}').decode())['ok'])
            bridge._window.create_file_dialog.assert_called_once()
