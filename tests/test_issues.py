"""Issue front matter (template fill and generated), banner logo and numbered Word outputs."""
import io
import json
from pathlib import Path
import tempfile
import unittest
import zipfile

from docx import Document
from docx.oxml.ns import qn

import app as application
import issue_export
from docx_export import generate_docx_from_form
from journal_templates import normalize_settings
from test_accounts import make_setup, register, headers
from test_formatter_templates import sample_article

ROOT = Path(__file__).resolve().parent.parent
PRESET = json.loads((ROOT / 'presets' / 'gastroia-journal-preset.json').read_text(encoding='utf-8'))
TEMPLATE = (ROOT / 'presets' / 'assets' / 'gastroia-jenerik-sablonu.docx').read_bytes()
ISSUE = {'volume': '9', 'issue': '2', 'month_tr': 'Ekim', 'month_en': 'October', 'year': '2025'}


def assets():
    result = {}
    for key, asset in PRESET['assets'].items():
        name, blob = (application.decode_template(asset) if key == 'jenerik' else application.decode_asset(asset)[:2])
        result[key] = (name, blob)
    return result


def all_text(document):
    return ' '.join(t.text or '' for t in document.element.iter(qn('w:t')))


class FrontMatterTests(unittest.TestCase):
    entries = issue_export.toc_entries([sample_article(), sample_article()], [(1, 15), (16, 16)])

    def test_toc_entries(self):
        self.assertEqual([e['pages'] for e in self.entries], ['1–15', '16'])
        self.assertEqual(self.entries[0]['authors'], 'Deniz Örnek, Alex Smith')

    def test_blank_template_is_filled_in_place(self):
        blob = issue_export.frontmatter_docx(ISSUE, PRESET['settings'], assets(), self.entries)
        document = Document(io.BytesIO(blob))
        text = ' '.join(all_text(document).split())
        header = ' '.join(t.text or '' for s in document.sections for t in s.header._element.iter(qn('w:t')))
        self.assertIn('E-ISSN : 2602-4144', text)
        self.assertIn('Ekim', header)
        self.assertIn('2025', header)
        self.assertIn('Akademik Yayıncılıkta Açık ve Tekrarlanabilir İş Akışları', text)
        self.assertIn('1 - 15', text)
        self.assertNotIn('Makalenin Türkçe Adı', text)
        self.assertEqual(text.count('Araştırma Makalesi / Research Article'), 2)
        self.assertIn('Prof. Dr. R. Cüneyt ERENOĞLU', text)  # imprint tables stay as designed
        self.assertEqual(len(list(document.element.body.iter(qn('w:sdt')))), 0)

    def test_generated_front_matter_without_template(self):
        given = {k: v for k, v in assets().items() if k != 'jenerik'}
        blob = issue_export.frontmatter_docx(ISSUE, PRESET['settings'], given, self.entries)
        document = Document(io.BytesIO(blob))
        text = all_text(document)
        self.assertGreaterEqual(len(document.sections), 2)
        self.assertIn('EDİTÖR KURLU', text)
        self.assertIn('E-ISSN: 2602-4144   Cilt | Volume: 9', text)
        self.assertIn('İÇİNDEKİLER', text)
        self.assertTrue(document.element.xpath('.//wp:anchor'))

    def test_imprint_parser(self):
        parsed = issue_export.parse_frontmatter('# A / B | 2\nX ; Y\n#| C\nZ\n---\nQ')
        self.assertEqual([s.get('title') for s in parsed], ['A / B', 'C', None, ''])
        self.assertTrue(parsed[1]['beside'])
        self.assertEqual(parsed[0]['entries'], [('X', 'Y')])


class BannerTests(unittest.TestCase):
    def test_banner_uses_full_text_width(self):
        settings = {**PRESET['settings'], 'logo_mode': 'banner', 'logo_width_cm': 0}
        data = sample_article()
        blob = generate_docx_from_form(data, {}, settings, {'logo': assets()['logo']})
        widths = [int(n.get('cx')) / 360000 for n in Document(io.BytesIO(blob)).element.xpath('.//wp:extent')]
        self.assertTrue(any(abs(w - 17.0) < 0.2 for w in widths), widths)  # 21 cm - 2 x 2 cm

    def test_explicit_logo_width_and_side_mode(self):
        settings = {**PRESET['settings'], 'logo_mode': 'side', 'logo_width_cm': 6}
        blob = generate_docx_from_form(sample_article(), {}, settings, {'logo': assets()['logo']})
        widths = [int(n.get('cx')) / 360000 for n in Document(io.BytesIO(blob)).element.xpath('.//wp:extent')]
        self.assertTrue(any(abs(w - 6.0) < 0.1 for w in widths), widths)
        with self.assertRaises(ValueError):
            normalize_settings({'logo_width_cm': 25})


def article_project(title, end):
    data = sample_article()
    data['cover']['tr_title'] = title
    data['cover']['start_page'], data['cover']['end_page'] = '1', str(end)
    data['sections'] = [{'id': '1', 'name': 'GİRİŞ', 'level': '1', 'content': 'Metin. ' * 30}]
    return {'format': 'aiditor-project', 'version': 1, 'data': data, 'figures': {}}


class IssueApiTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='aiditor-issue-test-')
        self.addCleanup(self.tmp.cleanup)
        app, _ = make_setup(Path(self.tmp.name))
        self.client = app.test_client()
        self.user = register(self.client)
        self.h = headers(self.user)
        saved = self.client.put('/api/journal', json={'settings': PRESET['settings'], 'assets': PRESET['assets'], 'base_revision': 0}, headers=self.h)
        self.assertEqual(saved.status_code, 200, saved.json)
        self.ids = []
        for title, end in (('Birinci makale', 1), ('İkinci makale', 16)):
            aid = application.uuid.uuid4().hex
            aid = str(application.uuid.UUID(aid))
            r = self.client.put('/api/articles/' + aid, json={'project': article_project(title, end), 'base_revision': 0}, headers=self.h)
            self.assertEqual(r.status_code, 200, r.json)
            self.ids.append(aid)
        self.issue = str(application.uuid.uuid4())
        data = {'volume': '9', 'issue': '2', 'month_tr': 'Ekim', 'month_en': 'October', 'year': '2025', 'first_page': '5',
                'articles': [{'id': self.ids[1]}, {'id': self.ids[0]}]}
        r = self.client.put('/api/issues/' + self.issue, json={'data': data, 'base_revision': 0}, headers=self.h)
        self.assertEqual(r.status_code, 200, r.json)

    def build(self, kind):
        return self.client.post(f'/api/issues/{self.issue}/build', json={'kind': kind}, headers=self.h)

    def test_issue_crud_and_validation(self):
        listing = self.client.get('/api/issues', headers=self.h).json
        self.assertEqual(listing['issues'][0]['title'], 'Cilt 9 Sayı 2 2025')
        bad = self.client.put('/api/issues/' + self.issue, json={'data': {'articles': [{'id': 'x'}], 'first_page': '1'}, 'base_revision': 1}, headers=self.h)
        self.assertEqual(bad.status_code, 400)
        conflict = self.client.put('/api/issues/' + self.issue, json={'data': {'first_page': '1', 'articles': []}, 'base_revision': 0}, headers=self.h)
        self.assertEqual(conflict.status_code, 409)

    def test_word_outputs_number_pages_consecutively(self):
        if True:
            response = self.build('articles_docx')
            self.assertEqual(response.status_code, 200, response.json)
            self.assertEqual([(r['start'], r['end']) for r in response.json['ranges']], [(5, 20), (21, 21)])
            blob = self.client.get('/download_file/' + response.json['key'], headers=self.h).data
            with zipfile.ZipFile(io.BytesIO(blob)) as archive:
                self.assertEqual(archive.namelist(), ['01-ikinci-makale.docx', '02-birinci-makale.docx'])
                first = Document(io.BytesIO(archive.read('01-ikinci-makale.docx')))
                self.assertIn('9', ' '.join(t.text or '' for t in first.element.iter(qn('w:t'))))
                self.assertIn('w:start="5"', first.element.xml.replace("'", '"') + first.sections[0]._sectPr.xml.replace("'", '"'))
            front = self.build('frontmatter_docx')
            self.assertEqual(front.status_code, 200, front.json)
            self.assertTrue(front.json['filename'].endswith('-jenerik.docx'))
            issue_zip = self.build('issue_docx')
            names = zipfile.ZipFile(io.BytesIO(self.client.get('/download_file/' + issue_zip.json['key'], headers=self.h).data)).namelist()
            self.assertEqual(names[0], '00-jenerik.docx')


if __name__ == '__main__':
    unittest.main()
