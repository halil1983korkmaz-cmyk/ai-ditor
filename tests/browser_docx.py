"""Saved page preferences and real editable Word downloads through both save paths."""
import base64
import io
import tempfile
from pathlib import Path
from docx import Document
from playwright.sync_api import sync_playwright, expect
from browser_workflow import Server, login, ready, saved, preset_saved, QA


def run():
    with tempfile.TemporaryDirectory(prefix='aiditor-word-browser-') as data, sync_playwright() as p:
        server=Server(data);browser=None
        try:
            url=server.start();browser=p.chromium.launch(headless=True);page=browser.new_page(viewport={'width':1360,'height':1000})
            errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
            login(page,url,'word_test',True,'Word Test Dergisi');ready(page)
            page.locator('#journal-tab').click()
            page.locator('#js-header_mode').select_option('odd_even')
            page.locator('#js-header_left').fill('TEK: {dergi}')
            page.locator('#js-header_even_left').fill('ÇİFT: {yazarlar}')
            page.locator('#js-header_first_mode').select_option('custom')
            page.locator('#js-header_first_center').fill('KAPAK: {yil}')
            page.locator('#js-footer_mode').select_option('same')
            page.locator('#js-footer_center').fill('Sayfa ')
            page.locator('[data-running-token="sayfa"]').click()
            expect(page.locator('#js-footer_center')).to_have_value('Sayfa {sayfa}')
            preset_saved(page);page.reload();ready(page);page.locator('#journal-tab').click()
            expect(page.locator('#js-header_mode')).to_have_value('odd_even')
            expect(page.locator('#js-header_even_left')).to_have_value('ÇİFT: {yazarlar}')
            page.locator('.running-settings').screenshot(path=str(QA/'running-settings.png'))
            page.locator('#articles-tab').click();page.locator('#c-tr-title').fill('Düzenlenebilir Word Denemesi')
            page.locator('#refs').fill('Örnek, A. (2026). Yayın. https://doi.org/10.1234/example')
            page.evaluate("""()=>{const field=document.querySelector('#sections-list [data-f=content]');if(!field)throw Error('Section editor missing');field.value='Örnek (2026), örnek bir kaynak sunmaktadır.';field.dispatchEvent(new Event('input',{bubbles:true}));}""")
            saved(page);page.locator('#btn-docx').click();expect(page.locator('#result-title')).to_have_text('Tamamlandı — Word belgesi hazır')
            expect(page.locator('#latex-next-steps')).to_be_hidden()
            expect(page.locator('#citation-result')).to_contain_text('1 metin içi atıf')
            with page.expect_download() as download:page.locator('#dl-docx').click()
            target=Path(data)/'browser.docx';download.value.save_as(target)
            doc=Document(target);assert doc.core_properties.title=='Düzenlenebilir Word Denemesi'
            assert doc.settings.odd_and_even_pages_header_footer
            assert 'ÇİFT' in doc.sections[0].even_page_header.tables[0].cell(0,0).text
            page.evaluate('''window.pywebview={api:{save_article_docx:async data=>{window.savedWord=data;return {ok:true,cancelled:false}}}}''')
            page.locator('#dl-docx').click();expect(page.locator('#docx-download-status')).to_have_text('Word dosyası kaydedildi.')
            assert base64.b64decode(page.evaluate('window.savedWord'))==target.read_bytes()
            page.route('**/download_docx/*',lambda route:route.fulfill(status=200,content_type='text/html',body='<html>Error</html>'))
            page.evaluate('window.savedWord=null');page.locator('#dl-docx').click()
            expect(page.locator('#docx-download-status')).to_contain_text('Word yerine farklı')
            assert page.evaluate('window.savedWord') is None
            page.set_viewport_size({'width':390,'height':844});page.locator('#journal-tab').click()
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth+1')
            assert not errors,errors
            print('PASS: saved odd/even/first-page settings, token insertion, Word browser/native bytes, error rejection, mobile layout')
        finally:
            if browser:browser.close()
            server.stop()

if __name__=='__main__':run()
