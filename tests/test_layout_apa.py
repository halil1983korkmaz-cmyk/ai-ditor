"""Editor-defined margins, paragraph spacing and APA 7 preferences (Word, LaTeX, checker, preset)."""
import copy
import io
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.shared import Cm

import app as application
from apa_rules import check_article, ordered_references
from docx_export import generate_docx_from_form
from formatter import generate_latex_from_form, turkish_sort_key
from journal_templates import default_settings, normalize_settings
from page_furniture import article_values
from test_accounts import make_setup, register, headers
from test_formatter_templates import sample_article

GASTROIA = {
    'template_id': 'scholarly', 'layout_mode': 'custom',
    'margin_top_cm': 0.58, 'margin_bottom_cm': 2, 'margin_left_cm': 2, 'margin_right_cm': 2,
    'header_distance_cm': 0.6, 'footer_distance_cm': 1.5, 'body_size': '11',
    'body_space_before_pt': 12, 'body_space_after_pt': 12, 'body_line_spacing': 1.15,
    'body_first_line_indent_cm': 0, 'ref_hanging_cm': 1.25, 'ref_space_before_pt': 6,
    'ref_space_after_pt': 0, 'ref_line_spacing': 1.0, 'heading_size_pt': 11,
    'heading_space_before_pt': 12, 'heading_space_after_pt': 6, 'caption_size_pt': 11, 'footnote_size_pt': 10,
}


def article():
    data = sample_article()
    data['sections'] = [{'id': '1', 'name': 'GİRİŞ', 'level': '1',
                         'content': 'Birinci paragraf metni. Kaya & Demir (2021) bunu göstermiştir.\n\nİkinci paragraf.'}]
    data['references'] = ('Zeybek, A. (2019). Son kayıt. Dergi, 1(1), 1–2.\n'
                          'Kaya, B. & Demir, C. (2021). Ortadaki kayıt. Dergi, 2(1), 3–4.\n'
                          'Aksoy, D. (2020). İlk kayıt. Dergi, 3(1), 5–6.')
    return data


def body_paragraphs(document):
    return [p for p in document.paragraphs if p.text.startswith(('Birinci paragraf', 'İkinci paragraf'))]


class SettingsTests(unittest.TestCase):
    def test_defaults_are_inert_and_complete(self):
        settings = default_settings()
        self.assertEqual(settings['layout_mode'], 'template')
        self.assertEqual((settings['apa_and'], settings['apa_et_al'], settings['apa_sort_references']), ('&', 'et al.', 'yes'))
        self.assertEqual(normalize_settings({})['margin_left_cm'], 2.5)

    def test_numeric_strings_from_the_form_and_bad_values(self):
        clean = normalize_settings({'margin_left_cm': '2,25', 'body_line_spacing': '1.15', 'ref_size_pt': ''})
        self.assertEqual((clean['margin_left_cm'], clean['body_line_spacing'], clean['ref_size_pt']), (2.25, 1.15, 0.0))
        for bad in ({'margin_top_cm': '9'}, {'margin_top_cm': 'abc'}, {'body_line_spacing': True},
                    {'layout_mode': 'free'}, {'apa_and': 'ile'}, {'apa_max_ref_authors': 0},
                    {'body_align': 'center'}, {'ref_size_pt': 3}):
            with self.assertRaises(ValueError, msg=bad):
                normalize_settings(bad)

    def test_apa_author_short_form_follows_preferences(self):
        data = {'cover': {}, 'authors': [{'name': 'A Karaca'}, {'name': 'B San'}]}
        self.assertEqual(article_values(data, normalize_settings({}))['yazarlar'], 'Karaca & San')
        self.assertEqual(article_values(data, normalize_settings({'apa_and': 've'}))['yazarlar'], 'Karaca ve San')
        data['authors'].append({'name': 'C Aad'})
        self.assertEqual(article_values(data, normalize_settings({}))['yazarlar'], 'Karaca et al.')
        self.assertEqual(article_values(data, normalize_settings({'apa_et_al': 'auto'}))['yazarlar'], 'Karaca vd.')
        self.assertEqual(article_values(data, normalize_settings({'apa_et_al': 'auto', 'english_only': True}))['yazarlar'], 'Karaca et al.')


class WordLayoutTests(unittest.TestCase):
    def build(self, **extra):
        return Document(io.BytesIO(generate_docx_from_form(article(), {}, {**GASTROIA, **extra}, {})))

    def test_template_mode_is_untouched(self):
        document = Document(io.BytesIO(generate_docx_from_form(article(), {}, {'template_id': 'scholarly'}, {})))
        self.assertAlmostEqual(document.sections[0].left_margin.cm, 2.5, places=1)
        self.assertEqual(body_paragraphs(document)[0].paragraph_format.line_spacing, 1.5)

    def test_custom_margins_reach_every_section(self):
        document = self.build()
        self.assertGreaterEqual(len(document.sections), 2)
        for section in document.sections:
            self.assertAlmostEqual(section.left_margin.cm, 2.0, places=1)
            self.assertAlmostEqual(section.right_margin.cm, 2.0, places=1)
        body = document.sections[-1]
        self.assertAlmostEqual(body.top_margin.cm, 0.58, places=1)
        self.assertAlmostEqual(body.bottom_margin.cm, 2.0, places=1)
        self.assertAlmostEqual(body.header_distance.cm, 0.6, places=1)
        self.assertAlmostEqual(body.footer_distance.cm, 1.5, places=1)

    def test_custom_paragraph_spacing_and_reference_indent(self):
        document = self.build()
        for paragraph in body_paragraphs(document):
            fmt = paragraph.paragraph_format
            self.assertEqual((fmt.space_before.pt, fmt.space_after.pt, fmt.line_spacing), (12, 12, 1.15))
            self.assertEqual(fmt.first_line_indent.cm, 0)
        refs = [p for p in document.paragraphs if p.text.startswith(('Aksoy', 'Kaya, B.', 'Zeybek'))]
        self.assertEqual([p.text.split(',')[0] for p in refs], ['Aksoy', 'Kaya', 'Zeybek'])
        for paragraph in refs:
            fmt = paragraph.paragraph_format
            self.assertAlmostEqual(fmt.left_indent.cm, 1.25, places=2)
            self.assertAlmostEqual(fmt.first_line_indent.cm, -1.25, places=2)
            self.assertEqual((fmt.space_before.pt, fmt.space_after.pt), (6, 0))
        heading = next(p for p in document.paragraphs if p.text == 'GİRİŞ')
        self.assertEqual((heading.paragraph_format.space_before.pt, heading.paragraph_format.space_after.pt), (12, 6))

    def test_reference_order_can_be_kept_as_entered(self):
        document = self.build(apa_sort_references='no')
        refs = [p.text.split(',')[0] for p in document.paragraphs if p.text.startswith(('Aksoy', 'Kaya, B.', 'Zeybek'))]
        self.assertEqual(refs, ['Zeybek', 'Kaya', 'Aksoy'])

    def test_first_page_cover_table_fits_custom_text_width(self):
        document = self.build(margin_left_cm=1.0, margin_right_cm=1.0)
        cover = document.tables[0]
        total = sum(column.width for column in cover.columns) / Cm(1)
        self.assertAlmostEqual(total, 19.0, places=1)


class LatexLayoutTests(unittest.TestCase):
    def test_custom_geometry_and_paragraph_format(self):
        tex = generate_latex_from_form(article(), {}, {**GASTROIA, 'template_id': 'classic'})
        self.assertIn(r'\geometry{a4paper,top=0.58cm,bottom=2cm,left=2cm,right=2cm,', tex)
        self.assertIn(r'\setlength{\parindent}{0cm}' + '\n' + r'\setlength{\parskip}{24pt}' + '\n' + r'\renewcommand{\baselinestretch}{1.15}', tex)
        self.assertIn(r'\setlength{\leftmargin}{1.25cm}', tex)
        self.assertIn(r'\setlength{\itemindent}{-1.25cm}', tex)

    def test_template_mode_keeps_original_geometry(self):
        tex = generate_latex_from_form(article(), {}, {'template_id': 'classic'})
        self.assertIn(r'\geometry{a4paper,top=1.5cm,bottom=1.5cm,left=1.5cm,right=1.5cm,', tex)


class ApaCheckTests(unittest.TestCase):
    def check(self, content, refs='', **settings):
        data = {'sections': [{'name': 'X', 'content': content}], 'references': refs, 'abstract': {}}
        report = check_article(data, normalize_settings(settings))
        return {item['code']: item for item in report['warnings']}

    def test_joiner_et_al_pages_and_dates(self):
        found = self.check('Kaya ve Demir (2021) ile Aad vd. (2012) görüştü (Sevindik, 2024, p. 45; Lyons, 2019, s. 4-6; Berger, n.d.).')
        self.assertEqual(found['citation_joiner']['examples'], ['Kaya ve Demir (2021'])
        self.assertEqual(found['citation_et_al']['examples'], ['Aad vd. (2012'])
        self.assertEqual(found['citation_pages']['examples'], ['p. 45', 's. 4-6'])
        self.assertEqual(found['no_date']['examples'], ['Berger, n.d.'])

    def test_english_manuscript_expects_english_forms(self):
        found = self.check('Aad vd. (2012) and (Karaca & San, 2021, s. 4).', english_only=True)
        self.assertIn('citation_et_al', found)
        self.assertIn('citation_pages', found)
        self.assertNotIn('citation_joiner', found)

    def test_configured_forms_pass(self):
        found = self.check('Karaca ve San (2021), Aad vd. (2012) (Sevindik, 2024, ss. 44-45; Berger, t.y.).', apa_and='ve', apa_et_al='vd.')
        self.assertEqual(set(found), set())

    def test_reference_rules(self):
        many = ', '.join(f'Yazar{chr(97 + i % 26)}{chr(97 + i // 26)}, A.' for i in range(22))
        refs = '\n'.join([f'{many} (2020). Çok yazarlı. Dergi, 1(1), 1–2.',
                          'Kaya, B. ve Demir, C. (2021). Bağlaç. Dergi, 1(1), 1–2. doi:10.1234/x',
                          'Yıl yok, A. Başlık. Dergi.'])
        found = self.check('Metin.', refs)
        self.assertEqual(set(found) & {'reference_authors', 'reference_joiner', 'doi_form', 'reference_year'},
                         {'reference_authors', 'reference_joiner', 'doi_form', 'reference_year'})
        self.assertIn('uncited_reference', found)

    def test_same_author_year_and_disabled(self):
        refs = 'Kaya, B. (2020). A.\nKaya, C. (2020). B.'
        self.assertIn('same_author_year', self.check('Kaya (2020).', refs))
        report = check_article({'sections': [], 'references': refs}, normalize_settings({'apa_check': 'off'}))
        self.assertEqual(report, {'enabled': False, 'warnings': []})

    def test_ordering_helper(self):
        data = {'references': 'Zeybek, A. (2019).\nAksoy, D. (2020).'}
        self.assertEqual(ordered_references(data, {}, turkish_sort_key)[0][:6], 'Aksoy,')
        self.assertEqual(ordered_references(data, {'apa_sort_references': 'no'}, turkish_sort_key)[0][:7], 'Zeybek,')


class PresetTests(unittest.TestCase):
    path = Path(__file__).resolve().parent.parent / 'presets' / 'gastroia-journal-preset.json'

    def test_preset_is_importable_and_matches_the_writing_guide(self):
        preset = json.loads(self.path.read_text(encoding='utf-8'))
        self.assertEqual((preset['format'], preset['version']), ('aiditor-journal-preset', 1))
        settings = normalize_settings(preset['settings'])
        self.assertEqual(settings, preset['settings'])
        self.assertEqual((settings['margin_top_cm'], settings['margin_bottom_cm'], settings['margin_left_cm'], settings['margin_right_cm']), (0.58, 2, 2, 2))
        self.assertEqual((settings['font_family'], settings['body_size'], settings['body_line_spacing']), ('texgyretermes', '11', 1.15))
        self.assertEqual((settings['body_space_before_pt'], settings['body_space_after_pt'], settings['footnote_size_pt']), (12, 12, 10))
        self.assertEqual((settings['apa_and'], settings['apa_et_al']), ('&', 'auto'))
        self.assertEqual(settings['issn_online'], '2602-4144')
        self.assertEqual(set(application.validate_assets(preset['assets'])), {'logo', 'license'})

    def test_preset_builds_a_word_document_and_matches_generator(self):
        preset = json.loads(self.path.read_text(encoding='utf-8'))
        assets = {}
        for key, asset in application.validate_assets(preset['assets']).items():
            name, blob, _ = application.decode_asset(asset)
            assets[key] = (name, blob)
        document = Document(io.BytesIO(generate_docx_from_form(article(), {}, preset['settings'], assets)))
        self.assertAlmostEqual(document.sections[-1].left_margin.cm, 2.0, places=1)
        self.assertIn('E-ISSN: 2602-4144', ' '.join(p.text for s in document.sections for p in s.first_page_header.paragraphs) +
                      ' '.join(c.text for s in document.sections for t in s.first_page_header.tables for r in t.rows for c in r.cells))


class ApiTests(unittest.TestCase):
    def test_apa_check_endpoint_and_saved_layout(self):
        with tempfile.TemporaryDirectory(prefix='aiditor-apa-test-') as tmp:
            app, _ = make_setup(Path(tmp))
            client = app.test_client()
            user = register(client)
            preset = json.loads(PresetTests.path.read_text(encoding='utf-8'))
            saved = client.put('/api/journal', json={'settings': preset['settings'], 'assets': preset['assets'], 'base_revision': 0}, headers=headers(user))
            self.assertEqual(saved.status_code, 200, saved.json)
            self.assertEqual(saved.json['settings']['margin_top_cm'], 0.58)
            data = article()
            data['sections'][0]['content'] = 'Kaya ve Demir (2021) bunu göstermiştir.'
            response = client.post('/api/apa_check', json={'data': data}, headers=headers(user))
            self.assertEqual(response.status_code, 200, response.json)
            self.assertIn('citation_joiner', {w['code'] for w in response.json['apa']['warnings']})
            bad = client.post('/api/journal', json={}, headers=headers(user))
            self.assertNotEqual(bad.status_code, 200)
            rejected = client.post('/validate_journal', json={'settings': {'margin_left_cm': 99}}, headers=headers(user))
            self.assertEqual(rejected.status_code, 400)


if __name__ == '__main__':
    unittest.main()
