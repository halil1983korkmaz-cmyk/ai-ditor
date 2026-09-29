"""Editable journal manuscripts built directly from the same form model as LaTeX."""
import copy
import io
import math
import re
import threading

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from formatter import (_cover_profile, _english_label, _normalize_table_model,
                       _numbered_section_title, _parse_table_rows, turkish_sort_key)
from journal_templates import normalize_settings
from citation_links import CitationIndex
from page_furniture import article_values, block_height_cm, citation_text, resolved_parts, running_slots

DOCX_MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
FONTS = {'texgyrepagella': 'Palatino Linotype', 'texgyretermes': 'Times New Roman',
         'texgyrebonum': 'Century Schoolbook', 'texgyreheros': 'Arial',
         'carlito': 'Calibri', 'latinmodern': 'Latin Modern Roman'}
ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
         'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}
_PDF_LOCK = threading.Lock()


def _element(tag: str, **attributes):
    element = OxmlElement('w:' + tag)
    for key, value in attributes.items():
        element.set(qn('w:' + key), str(value))
    return element


def _paragraph(container, text='', *, size=None, bold=False, italic=False, color=None,
               align='left', before=0, after=3, keep=False, style=None):
    p = container.add_paragraph(style=style)
    p.alignment = ALIGN[align]
    fmt = p.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    fmt.line_spacing = 1.08
    fmt.keep_with_next = keep
    fmt.widow_control = True
    _text(p, str(text), size=size, bold=bold, italic=italic, color=color)
    return p


def _text(p, text, *, size=None, bold=False, italic=False, color=None):
    """Keep plain prose and URLs editable; URLs also receive Word relationships."""
    position = 0
    for match in re.finditer(r'https?://[^\s<>]+', text):
        _run(p, text[position:match.start()], size, bold, italic, color)
        url = match.group(0).rstrip('.,;)')
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
        run = _run(p, url, size, bold, italic, color or '245F91')
        link.append(run._r)
        p._p.append(link)
        position = match.start() + len(url)
    _run(p, text[position:], size, bold, italic, color)


def _run(p, text, size=None, bold=False, italic=False, color=None):
    run = p.add_run(text)
    if size:
        run.font.size = Pt(float(size))
    run.bold, run.italic = bold, italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color.lstrip('#'))
    return run


def _cell_style(cell, *, fill=None, border='none', padding=70):
    props = cell._tc.get_or_add_tcPr()
    for old in list(props):
        if old.tag in {qn('w:tcBorders'), qn('w:tcMar'), qn('w:shd')}:
            props.remove(old)
    borders = _element('tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        borders.append(_element(edge, val='single' if border != 'none' else 'nil', sz=4, color=border if border != 'none' else 'FFFFFF'))
    props.append(borders)
    margins = _element('tcMar')
    for edge in ('top', 'left', 'bottom', 'right'):
        margins.append(_element(edge, w=padding, type='dxa'))
    props.append(margins)
    if fill:
        props.append(_element('shd', fill=fill.lstrip('#'), val='clear'))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _table(container, widths, rows=1):
    if hasattr(container, 'sections') or hasattr(container, '_tc'):
        table = container.add_table(rows=rows, cols=len(widths))
    else:
        table = container.add_table(rows=rows, cols=len(widths), width=Cm(sum(widths)))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for old in list(table._tbl.tblPr.findall(qn('w:tblW'))):
        table._tbl.tblPr.remove(old)
    table._tbl.tblPr.append(_element('tblW', w=round(sum(widths) * 1440 / 2.54), type='dxa'))
    for index, width in enumerate(widths):
        table.columns[index].width = Cm(width)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Cm(widths[index])
            _cell_style(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1
            p.add_run().font.size = Pt(1)
    return table


def _compact_cell(cell):
    """Remove only an unused leading placeholder once real block content exists."""
    if len(cell.paragraphs) > 1 and not cell.paragraphs[0].text and not cell.paragraphs[0]._p.xpath('.//w:drawing'):
        cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)


def _picture(p, asset, *, max_width_cm=18, max_height_cm=14):
    filename, blob = asset
    if filename.lower().endswith('.pdf'):
        import pypdfium2 as pdfium
        with _PDF_LOCK:
            pdf = pdfium.PdfDocument(blob)
            try:
                page = pdf[0]
                try:
                    scale = min(3, 2800 / max(page.get_size()))
                    bitmap = page.render(scale=scale)
                    try:
                        image = bitmap.to_pil()
                        stream = io.BytesIO()
                        image.save(stream, format='PNG')
                        blob = stream.getvalue()
                    finally:
                        bitmap.close()
                finally:
                    page.close()
            finally:
                pdf.close()
    with Image.open(io.BytesIO(blob)) as image:
        width, height = image.size
        # Normalize formats, orientation and CMYK for Word and LibreOffice.
        from PIL import ImageOps
        image = ImageOps.exif_transpose(image)
        image = image.convert('RGBA' if 'A' in image.getbands() else 'RGB')
        width, height = image.size
        stream = io.BytesIO()
        image.save(stream, format='PNG')
    ratio = min(max_width_cm / width, max_height_cm / height)
    shape = p.add_run().add_picture(io.BytesIO(stream.getvalue()), width=Cm(width * ratio), height=Cm(height * ratio))
    shape._inline.docPr.set('descr', filename)


def _page_field(p, value='1', size=8):
    # Separate runs are required by Word and LibreOffice's complex-field reader.
    p.add_run()._r.append(_element('fldChar', fldCharType='begin'))
    code = _element('instrText')
    code.set(qn('xml:space'), 'preserve')
    code.text = ' PAGE '
    p.add_run()._r.append(code)
    p.add_run()._r.append(_element('fldChar', fldCharType='separate'))
    p.add_run(value).font.size = Pt(size)
    p.add_run()._r.append(_element('fldChar', fldCharType='end'))


def _running_part(part, slots, values, *, size, line, kind, start, width=18, scholarly=False):
    part.is_linked_to_previous = False
    p0 = part.paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.line_spacing = Pt(1)
    p0.add_run().font.size = Pt(1)
    if not any(slots):
        return
    only_center = not slots[0] and not slots[2]
    widths = [width * .88, width * .01, width * .11] if scholarly and not slots[1] and slots[0] else [width / 3] * 3
    table = _table(part, [width] if only_center else widths)
    entries = [(table.cell(0, 0), slots[1], 'center')] if only_center else [(table.cell(0, i), text, alignment) for i, (text, alignment) in enumerate(zip(slots, ('left', 'center', 'right')))]
    for cell, text, alignment in entries:
        p = cell.paragraphs[0]
        p.alignment = ALIGN[alignment]
        p.paragraph_format.line_spacing = 1.1
        for kind_, value in resolved_parts(text, values):
            if kind_ == 'page':
                _page_field(p, start, size)
            else:
                _text(p, value, size=size, bold=scholarly and kind == 'header', italic=scholarly and kind == 'header')
        if scholarly and kind == 'header':
            for run in p.runs:
                run.bold = run.italic = True
        if line == 'line':
            borders = _element('tcBorders')
            borders.append(_element('bottom' if kind == 'header' else 'top', val='single', sz=4, color='777777'))
            props = cell._tc.get_or_add_tcPr()
            for old in list(props.findall(qn('w:tcBorders'))):
                props.remove(old)
            props.append(borders)


def _setup(doc, settings, data):
    section = doc.sections[0]
    scholarly = settings['template_id'] == 'scholarly'
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.5 if scholarly else 1.5)
    section.header_distance = section.footer_distance = Cm(.8)
    section.top_margin = Cm(.8 + block_height_cm(settings, data, 'header') + .4)
    section.bottom_margin = Cm(.8 + block_height_cm(settings, data, 'footer') + .4)
    if scholarly:
        section.header_distance = Cm(1.25)
        section.top_margin = Cm(1.25 + block_height_cm(settings, data, 'header') + .65)
        section.bottom_margin = Cm(max(2.5, section.bottom_margin.cm))
    section.different_first_page_header_footer = True
    doc.settings.odd_and_even_pages_header_footer = any(settings[kind + '_mode'] == 'odd_even' for kind in ('header', 'footer'))
    start = str(data.get('cover', {}).get('start_page', '1'))
    start = start if start.isdigit() else '1'
    section._sectPr.insert_element_before(_element('pgNumType', start=start), 'w:cols', 'w:titlePg', 'w:docGrid')
    values = article_values(data, settings)
    for kind in ('header', 'footer'):
        for variant, attr in [('odd', kind), ('even', 'even_page_' + kind), ('first', 'first_page_' + kind)]:
            _running_part(getattr(section, attr), running_slots(settings, data, kind, variant), values,
                          size=int(settings[kind + '_font_size']), line=settings[kind + '_rule'], kind=kind, start=start,
                          width=16 if scholarly else 18, scholarly=scholarly)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption', 'Header', 'Footer'):
        style = doc.styles[name]
        style.font.name = FONTS[settings['font_family']]
        for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            style.element.rPr.rFonts.attrib.pop(qn('w:' + attr), None)
        if style.element.pPr is not None:
            for border in list(style.element.pPr.findall(qn('w:pBdr'))):
                style.element.pPr.remove(border)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(int(settings['body_size']))
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.08
    for name in ('Heading 1', 'Heading 2', 'Heading 3'):
        style = doc.styles[name]
        style.font.bold = True
        style.font.size = Pt(11)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
    doc.core_properties.title = values['baslik']
    doc.core_properties.author = '; '.join(a.get('name', '') for a in data.get('authors', []))
    doc.core_properties.subject = values['dergi']
    doc.core_properties.comments = ''


def _start_cover_footer(doc):
    """Put editorial notes above any user-selected first-page running footer."""
    footer = doc.sections[0].first_page_footer
    original = list(footer._element)
    for child in original:
        footer._element.remove(child)
    return footer, original


def _finish_cover_footer(doc, footer, original, settings, data):
    # Word footers grow upward from the bottom margin. Reserve their text/image
    # height so the final abstract line cannot overlap them during pagination.
    paragraphs = list(footer.paragraphs)
    section = doc.sections[0]
    width = (section.page_width - section.left_margin - section.right_margin) / Cm(1)
    height = 0.0
    for paragraph in paragraphs:
        fmt = paragraph.paragraph_format
        sizes = [run.font.size.pt for run in paragraph.runs if run.font.size]
        size = max(sizes, default=8)
        lines = sum(max(1, math.ceil(len(line) / max(1, width * 28.346 / (size * .55)))) for line in paragraph.text.split('\n'))
        spacing = fmt.line_spacing
        leading = spacing.pt if hasattr(spacing, 'pt') else size * 1.2 * (spacing or 1.08)
        height += (lines * leading + (fmt.space_before.pt if fmt.space_before else 0) + (fmt.space_after.pt if fmt.space_after else 0)) * 2.54 / 72
        height += sum(int(node.get('cy')) / Cm(1) for node in paragraph._p.xpath('.//wp:extent'))
    if paragraphs and settings['footer_layout'] != 'minimal':
        border = _element('pBdr')
        border.append(_element('top', val='single', sz=4, color=settings['accent_color'].lstrip('#')))
        paragraphs[0]._p.get_or_add_pPr().append(border)
    for child in original:
        footer._element.append(child)
    if any(running_slots(settings, data, 'footer', 'first')):
        height += block_height_cm(settings, data, 'footer')
    reserved = max(1.85, section.footer_distance.cm + height + .45)
    if reserved > 20:
        raise ValueError('Kapak dipnotları bir sayfaya sığmayacak kadar uzun. Uzun açıklamaları makalenin beyan bölümlerine taşıyın.')
    section.bottom_margin = Cm(reserved)


def _body_section(doc, settings, data):
    """Continue the issue numbering with normal odd/even headers after the cover."""
    from docx.enum.section import WD_SECTION_START
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.different_first_page_header_footer = False
    scholarly = settings['template_id'] == 'scholarly'
    section.header_distance = Cm(1.25 if scholarly else .8)
    section.top_margin = Cm((1.25 if scholarly else .8) + block_height_cm(settings, data, 'header') + (.65 if scholarly else .4))
    section.bottom_margin = Cm(max(2.5 if scholarly else 1.85, .8 + block_height_cm(settings, data, 'footer') + .4))
    # The cloned section must continue numbering, not restart at the article's first page.
    for number in list(section._sectPr.findall(qn('w:pgNumType'))):
        section._sectPr.remove(number)


def _scholarly_cover(doc, data, settings, assets):
    cov, abstract = data.get('cover', {}), data.get('abstract', {})
    english = settings['english_only']
    scale = float(_cover_profile(data, settings)['abstract'][0]) / 9
    values = article_values(data, settings)
    section = doc.sections[0]
    section.header_distance = Cm(.3)
    section.top_margin = Cm(.3 + block_height_cm(settings, data, 'header') + .25)
    section.bottom_margin = Cm(2)
    accent = settings['accent_color'].lstrip('#')
    table = _table(doc, [3.1, 12.9])
    left, right = table.rows[0].cells
    if assets.get('logo') and settings['show_logo']:
        _picture(left.paragraphs[0], assets['logo'], max_width_cm=2.9, max_height_cm=settings['logo_height_cm'])
    _cell_style(right, fill='F2F2F2', padding=100)
    names = [values['dergi']]
    if not english and settings['journal_name_en'] and settings['journal_name_en'] != values['dergi']:
        names.append(settings['journal_name_en'])
    _paragraph(right, ' / '.join(filter(None, names)), size=15, bold=True, italic=True, align='center', after=8)
    for labels in ([('Year', 'Volume', 'Issue')] if english else [('Yıl', 'Cilt', 'Sayı'), ('Year', 'Volume', 'Issue')]):
        text = '  '.join(label + ': ' + str(cov[key]) for label, key in zip(labels, ('year', 'volume', 'issue')) if cov.get(key))
        if text:
            _paragraph(right, text, size=11, italic=True, align='center', after=0)
    details = [label + ': ' + settings[key] for label, key in [('ISSN', 'issn_print'), ('e-ISSN', 'issn_online')] if settings[key]]
    if settings['journal_url']:
        details.append(settings['journal_url'])
    doi = str(cov.get('doi', '')).removeprefix('https://doi.org/').removeprefix('doi:').strip()
    if doi and settings['doi_position'] == 'top':
        details.append('https://doi.org/' + doi)
    for text in details:
        _paragraph(right, text, size=9, align='center', before=4)
    _compact_cell(right)
    rule = _paragraph(doc, after=14, size=1)
    border = _element('pBdr');border.append(_element('bottom', val='single', sz=16, color=accent))
    rule._p.get_or_add_pPr().append(border)
    if cov.get('article_type'):
        p = _paragraph(doc, after=0)
        _run(p, 'Article Type: ' if english else 'Makale Türü / Article Type: ', 10, bold=True)
        _text(p, _english_label(cov['article_type']) if english else cov['article_type'], size=10, italic=True)
    if settings['footer_layout'] != 'minimal':
        p = _paragraph(doc, after=2, align='justify')
        _run(p, 'Citation: ' if english else 'Atıf / Citation: ', 10, bold=True)
        _text(p, citation_text(data, settings), size=10)
    has_tr = bool(abstract.get('tr_abs', '').strip()) and not english
    has_en = bool(abstract.get('en_abs', '').strip())
    english_cover = english or (has_en and not has_tr)
    title = (cov.get('en_title') or cov.get('tr_title', '')) if english_cover else (cov.get('tr_title') or cov.get('en_title', ''))
    p = _paragraph(doc, title + (' *' if cov.get('title_note') else ''), size=14 * scale, bold=True,
                   align='center', before=18, after=8, keep=True, style='Title')
    p.paragraph_format.line_spacing = Pt(23 * scale)
    if not english and not has_en and cov.get('en_title') and cov['en_title'] != cov.get('tr_title'):
        _paragraph(doc, cov['en_title'], size=11, italic=True, align='center', after=8, keep=True, style='Subtitle')
    authors = [a for a in data.get('authors', []) if a.get('name', '').strip()]
    if authors:
        p = _paragraph(doc, after=4, keep=True)
        for index, author in enumerate(authors):
            if index:
                _run(p, ', ', 11)
            _run(p, author['name'], 11)
            _run(p, str(index + 1), 7).font.superscript = True
        for index, author in enumerate(authors):
            info = ' · '.join(str(author.get(key, '')).strip() for key in ('title', 'affiliation') if str(author.get(key, '')).strip())
            orcid = str(author.get('orcid', '')).strip().removeprefix('https://orcid.org/')
            if orcid:
                info += (' · ' if info else '') + 'https://orcid.org/' + orcid
            if author.get('email') and not author.get('corresponding'):
                info += (' · ' if info else '') + author['email']
            if info:
                _paragraph(doc, str(index + 1) + ' ' + info, size=9, italic=True, after=4)

    def add_abstract(language, label, keyword_label):
        size = 10 * scale if language == 'tr' or english_cover else 10
        _paragraph(doc, label, size=10, bold=True, align='center', before=12, after=8, keep=True)
        for text in re.split(r'\n\s*\n', abstract.get(language + '_abs', '').strip()):
            if text:
                p = _paragraph(doc, text, size=size, align='justify', after=0)
                p.paragraph_format.first_line_indent = Cm(1.25)
                p.paragraph_format.line_spacing = 1.5
        if abstract.get(language + '_kw'):
            p = _paragraph(doc, before=5, after=4)
            p.paragraph_format.first_line_indent = Cm(1.25)
            _run(p, keyword_label + ': ', 10, bold=True, italic=True)
            _text(p, abstract[language + '_kw'], size=10)
    if english_cover and has_en:
        add_abstract('en', settings['english_abstract_heading'], 'Keywords')
    elif has_tr:
        add_abstract('tr', 'Öz', 'Anahtar kelimeler')
    footer, original = _start_cover_footer(doc)
    if doi and settings['doi_position'] == 'bottom':
        _paragraph(footer, 'https://doi.org/' + doi, size=8, after=2)
    for label, value in [('Ethics Statement: ' if english else 'Etik Beyan: ', cov.get('ethics')),
                         ('* ', cov.get('title_note')), ('', settings['footer_text']),
                         ('Editor: ' if english else 'Editör / Editor: ', cov.get('editor'))]:
        if value:
            _paragraph(footer, label + value, size=8, before=4)
    if assets.get('license') and settings['show_cc_logo']:
        _picture(_paragraph(footer), assets['license'], max_width_cm=5, max_height_cm=.5)
    # Editorial contact/dates belong to the cover footer; the user-controlled
    # running first-page footer, if present, remains in the same editable part.
    contacts = []
    for author in authors:
        if author.get('corresponding'):
            contacts.append(' · '.join(str(author.get(k, '')).strip() for k in ('name', 'email', 'orcid') if str(author.get(k, '')).strip()))
    p = None
    if contacts:
        p = _paragraph(footer, size=8, after=0)
        _run(p, 'Corresponding Author: ' if english else 'Sorumlu Yazar / Corresponding Author: ', 8, bold=True)
        _text(p, '; '.join(contacts), size=8)
    labels = [('Received', 'received'), ('Accepted', 'accepted'), ('Published', 'published')] if english else [('Gönderim / Received', 'received'), ('Kabul / Accepted', 'accepted'), ('Yayımlanma / Published', 'published')]
    dates = [label + ': ' + str(cov[key]) for label, key in labels if cov.get(key)]
    if dates:
        dates_p = _paragraph(footer, '    '.join(dates), size=8, after=0)
        p = p or dates_p
    _finish_cover_footer(doc, footer, original, settings, data)
    if has_tr and has_en and not english:
        _body_section(doc, settings, data)
        p = _paragraph(doc, cov.get('en_title') or title, size=14, bold=True, align='center', after=16, keep=True, style='Subtitle')
        p.paragraph_format.line_spacing = Pt(23)
        add_abstract('en', settings['english_abstract_heading'], 'Keywords')


def _cover(doc, data, settings, assets):
    if settings['template_id'] == 'scholarly':
        _scholarly_cover(doc, data, settings, assets)
        return
    cov = data.get('cover', {})
    english = settings['english_only']
    layout = settings['template_id']
    accent = settings['accent_color'].lstrip('#')
    profile = _cover_profile(data, settings)
    values = article_values(data, settings)
    def identity(container, align='left', color=None):
        _paragraph(container, values['dergi'], size=11, bold=True, align=align, color=color)
        secondary = settings['journal_name_en']
        if not english and secondary and secondary != values['dergi']:
            _paragraph(container, secondary, size=9, italic=True, align=align, color=color)
    def details(container, align='left'):
        identifiers = [label + ': ' + settings[key] for label, key in [('ISSN', 'issn_print'), ('e-ISSN', 'issn_online')] if settings[key]]
        if identifiers:
            _paragraph(container, '  ·  '.join(identifiers), size=7.5, align=align)
        if settings['journal_url']:
            _paragraph(container, settings['journal_url'], size=7.5, align=align)
        if settings['doi_position'] == 'top' and cov.get('doi'):
            _paragraph(container, 'DOI: https://doi.org/' + cov['doi'], size=7.5, align=align)
    def logo(container, align='left'):
        if assets.get('logo') and settings['show_logo']:
            p = _paragraph(container, align=align, after=3)
            _picture(p, assets['logo'], max_width_cm=4.2, max_height_cm=settings['logo_height_cm'])
    if layout == 'centered':
        logo(doc, 'center'); identity(doc, 'center'); details(doc, 'center')
    elif layout == 'contemporary':
        table = _table(doc, [13, 5])
        left, right = table.rows[0].cells
        _cell_style(left, fill=accent, padding=110)
        identity(left, color='FFFFFF'); logo(right, 'right')
        _compact_cell(left); _compact_cell(right)
        details(doc)
    else:
        table = _table(doc, [5, 13] if layout == 'classic' else [13, 5])
        left, right = table.rows[0].cells
        if layout == 'classic':
            logo(left); identity(right, 'right'); details(right, 'right')
        else:
            identity(left); details(left); logo(right, 'right')
        _compact_cell(left); _compact_cell(right)
    meta = '  ·  '.join(label + ': ' + cov[key] for label, key in [('Year' if english else 'Yıl', 'year'), ('Volume' if english else 'Cilt', 'volume'), ('Issue' if english else 'Sayı', 'issue')] if cov.get(key))
    if meta:
        _paragraph(doc, meta, size=8, before=6, after=6, align='center' if layout == 'centered' else 'left')
    if cov.get('article_type'):
        _paragraph(doc, _english_label(cov['article_type']) if english else cov['article_type'], size=8, italic=True, after=4)
    titles = [(values['baslik'], 'title', False)]
    if not english and cov.get('en_title') and cov['en_title'] != values['baslik']:
        titles.append((cov['en_title'], 'subtitle', True))
    for text, key, italic in titles:
        if key == 'title' and cov.get('title_note'):
            text += ' *'
        _paragraph(doc, text, size=profile[key][0], bold=True, italic=italic, keep=True,
                   color=accent if layout == 'contemporary' else None,
                   align='center' if layout == 'centered' else 'left', after=4,
                   style='Title' if key == 'title' else 'Subtitle')
    authors = [a for a in data.get('authors', []) if a.get('name', '').strip()]
    if authors:
        p = _paragraph(doc, align='center' if layout == 'centered' else 'left', after=6, keep=True)
        for i, author in enumerate(authors):
            if i:
                _run(p, ', ', profile['authors'][0])
            _run(p, author['name'], profile['authors'][0], bold=True)
            r = _run(p, str(i + 1) + (settings['corresponding_marker'] if author.get('corresponding') else ''), 7)
            r.font.superscript = True
    date_items = [(('Received' if english else 'Başvuru / Received'), cov.get('received')),
                  (('Accepted' if english else 'Kabul / Accepted'), cov.get('accepted')),
                  (('Published' if english else 'Yayın / Published'), cov.get('published')),
                  (('Editor' if english else 'Editör / Editor'), cov.get('editor'))]
    abstract = data.get('abstract', {})
    languages = ([] if english or not abstract.get('tr_abs') else [('Özet', 'tr', 'Anahtar kelimeler')])
    if abstract.get('en_abs'):
        languages.append((settings['english_abstract_heading'], 'en', 'Keywords'))
    def add_abstract(container, item):
        label, lang, keywords = item
        _paragraph(container, label, size=9, bold=True, color=accent, before=3, keep=True)
        _paragraph(container, abstract[lang + '_abs'], size=profile['abstract'][0], italic=True, align='justify')
        if abstract.get(lang + '_kw'):
            p = _paragraph(container, size=7.5, after=5)
            _run(p, keywords + ': ', 7.5, bold=True)
            _text(p, abstract[lang + '_kw'], size=7.5)
    if layout == 'classic' and languages:
        table = _table(doc, [3.5, 14.5])
        left, right = table.rows[0].cells
        for label, value in date_items:
            if value:
                _paragraph(left, label, size=profile['info'][0], bold=True, keep=True)
                _paragraph(left, value, size=profile['info'][0], after=6)
        for item in languages:
            add_abstract(right, item)
        _compact_cell(left); _compact_cell(right)
    else:
        dates = '  ·  '.join(label + ': ' + value for label, value in date_items if value)
        if dates:
            _paragraph(doc, dates, size=7.5, after=5)
        if layout == 'contemporary' and len(languages) == 2:
            table = _table(doc, [9, 9])
            for cell, item in zip(table.rows[0].cells, languages):
                add_abstract(cell, item); _compact_cell(cell)
        else:
            for item in languages:
                add_abstract(doc, item)
    footer, original = _start_cover_footer(doc)
    fs = float(profile['footer'][0])
    for i, author in enumerate(authors):
        info = ' · '.join(str(author.get(k, '')).strip() for k in ('title', 'affiliation', 'email') if str(author.get(k, '')).strip())
        orcid = str(author.get('orcid', '')).strip().removeprefix('https://orcid.org/')
        if orcid:
            info += (' · ' if info else '') + 'https://orcid.org/' + orcid
        _paragraph(footer, str(i + 1) + ' ' + author['name'] + (': ' + info if info else ''), size=fs, after=2)
    corresponding = [a for a in authors if a.get('corresponding')]
    if corresponding:
        _paragraph(footer, settings['corresponding_marker'] + ' ' + ('Corresponding author: ' if english else 'Sorumlu yazar / Corresponding author: ') + '; '.join(a['name'] + (' · ' + a['email'] if a.get('email') else '') for a in corresponding), size=fs, italic=True)
    if settings['footer_layout'] != 'minimal':
        _paragraph(footer, ('Suggested Citation: ' if english else 'Önerilen Atıf / Suggested Citation: ') + citation_text(data, settings), size=fs)
    for label, value in [('Ethics Statement: ' if english else 'Etik Beyan / Ethics Statement: ', cov.get('ethics')), ('* ', cov.get('title_note')), ('', settings['footer_text'])]:
        if value:
            _paragraph(footer, label + value, size=fs)
    if settings['doi_position'] == 'bottom' and cov.get('doi'):
        _paragraph(footer, 'DOI: https://doi.org/' + cov['doi'], size=fs)
    if assets.get('license') and settings['show_cc_logo']:
        p = _paragraph(footer, after=2)
        _picture(p, assets['license'], max_width_cm=6, max_height_cm=.55)
    _finish_cover_footer(doc, footer, original, settings, data)


def _figure_or_table(doc, item, figures, english):
    section = doc.sections[0]
    available_width = (section.page_width - section.left_margin - section.right_margin) / Cm(1)
    caption = (item.get('en_cap') or item.get('tr_cap', '')) if english else ' / '.join(t for t in (item.get('tr_cap'), item.get('en_cap')) if t)
    is_figure = item['type'] == 'figure'
    label = ('Figure' if english else 'Şekil') if is_figure else ('Table' if english else 'Tablo')
    title = label + ' ' + str(item.get('number', '1')) + ('. ' + caption if caption else '')
    if is_figure:
        asset = figures.get(item.get('file_key'))
        if asset is None:
            raise ValueError('Word çıktısı için şekil dosyası eksik.')
        try:
            width = float(item.get('fig_width', 90))
        except (TypeError, ValueError):
            width = 90
        width = max(50, min(100, width)) if math.isfinite(width) else 90
        p = _paragraph(doc, align='center', keep=True, before=4)
        _picture(p, asset, max_width_cm=available_width * width / 100, max_height_cm=13)
        _paragraph(doc, title, size=9, bold=True, align='center', after=6, style='Caption')
        return
    _paragraph(doc, title, size=9, bold=True, align='center', keep=True, before=4, style='Caption')
    source = copy.deepcopy(item)
    if not isinstance(source.get('tbl_model'), dict) or not source['tbl_model'].get('rows'):
        source['tbl_model'] = {'rows': [[{'text': cell} for cell in row] for row in _parse_table_rows(source.get('tbl_data', ''))], 'header_rows': 1}
    model = _normalize_table_model(source)
    table = _table(doc, [available_width * width for width in model['widths']], rows=len(model['rows']))
    try:
        fs = float(item.get('tbl_fontsize') or 9)
    except (TypeError, ValueError):
        fs = 9
    fs = max(7, min(12, fs)) if math.isfinite(fs) else 9
    for row_index, entries in enumerate(model['layout']):
        row = table.rows[row_index]
        if row_index < model['header_rows']:
            row._tr.get_or_add_trPr().append(_element('tblHeader', val='true'))
        for entry in entries:
            col = entry['col']
            cell = table.cell(row_index, col)
            end_row = min(len(model['rows']) - 1, row_index + entry['rowspan'] - 1)
            end_col = min(model['ncols'] - 1, col + entry['colspan'] - 1)
            if end_row != row_index or end_col != col:
                cell = cell.merge(table.cell(end_row, end_col))
            value = entry['cell']
            cell.text = ''
            _cell_style(cell, fill=value.get('bgcolor') or ('E9EEEF' if row_index < model['header_rows'] else None),
                        border='B7C3C7' if model['grid_borders'] else 'none', padding=70)
            p = cell.paragraphs[0]
            p.alignment = ALIGN.get(value.get('align'), WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.08
            _text(p, value['text'], size=fs, bold=value.get('bold') or row_index < model['header_rows'], italic=value.get('italic', False), color=value.get('textcolor') or None)
            if value.get('underline'):
                for run in p.runs:
                    run.underline = True
    _paragraph(doc, after=3, size=2)


def _body(doc, data, settings, figures):
    english = settings['english_only']
    scholarly = settings['template_id'] == 'scholarly'
    refs = sorted([r.strip() for r in data.get('references', '').splitlines() if r.strip()], key=turkish_sort_key)
    citations = CitationIndex(refs, settings['link_citations'])
    items = data.get('figtables', [])
    placed = set()
    for section in data.get('sections', []):
        title = section.get('name', '').strip()
        detected = _numbered_section_title(title)
        if detected:
            title = detected[0]
        title = (_english_label(title) if english else title) or ('Section' if english else 'Bölüm')
        level = max(1, min(3, int(section.get('level', '1'))))
        _paragraph(doc, title, style='Heading ' + str(level), bold=True, keep=True, before=8, after=4,
                   align='center' if level == 1 and not scholarly else 'left')
        assigned = [(i, item) for i, item in enumerate(items) if i not in placed and
                    (str(item.get('section_id')) == str(section.get('id')) if item.get('section_id') is not None and section.get('id') is not None else item.get('section', '').strip() == section.get('name', '').strip())]
        for i, item in assigned:
            if item.get('placement') == 'section_start':
                _figure_or_table(doc, item, figures, english); placed.add(i)
        for para in re.split(r'\n\s*\n', section.get('content', '').strip()):
            if para.strip():
                p = _paragraph(doc, align='justify')
                if scholarly:
                    p.paragraph_format.first_line_indent = Cm(1.25)
                    p.paragraph_format.line_spacing = 1.5
                    p.paragraph_format.space_after = Pt(0)
                cursor = 0
                for start, end, index in citations.spans(para.strip()):
                    _text(p, para.strip()[cursor:start])
                    link = _element('hyperlink', anchor='aiditor_ref_' + str(index), history='1')
                    run = _run(p, para.strip()[start:end])
                    link.append(run._r)
                    p._p.append(link)
                    cursor = end
                _text(p, para.strip()[cursor:])
                for i, item in assigned:
                    anchor = item.get('after_para', '').strip()
                    if i not in placed and anchor and anchor.lower() in para.lower():
                        _figure_or_table(doc, item, figures, english); placed.add(i)
        for i, item in assigned:
            if i not in placed:
                _figure_or_table(doc, item, figures, english); placed.add(i)
    for i, item in enumerate(items):
        if i not in placed:
            _figure_or_table(doc, item, figures, english)
    labels = [('ack', 'Acknowledgements' if english else 'Teşekkür / Acknowledgements'),
              ('contrib', 'Author Contributions' if english else 'Araştırmacıların Katkı Oranı / Author Contributions'),
              ('conflict', 'Conflict of Interest' if english else 'Çıkar Çatışması / Conflict of Interest')]
    for key, title in labels:
        text = data.get('extra', {}).get(key, '').strip()
        if text:
            _paragraph(doc, title, style='Heading 1', bold=True, keep=True, before=8)
            _paragraph(doc, text, align='justify')
    refs = sorted([ref.strip() for ref in data.get('references', '').splitlines() if ref.strip()], key=turkish_sort_key)
    if refs:
        _paragraph(doc, 'References' if english else ('Kaynaklar' if scholarly else 'Kaynakça / References'), style='Heading 1', bold=True, keep=True, before=8, align='center' if scholarly else 'left')
        for index, ref in enumerate(refs):
            p = _paragraph(doc, ref)
            p._p.insert(1, _element('bookmarkStart', id=index, name='aiditor_ref_' + str(index)))
            p._p.append(_element('bookmarkEnd', id=index))
            p.paragraph_format.left_indent = Cm(.5)
            p.paragraph_format.first_line_indent = Cm(-.5)
            if scholarly:
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_after = Pt(6)
                for run in p.runs:
                    run.font.size = Pt(10)


def generate_docx_from_form(data: dict, figures: dict, journal_settings: dict | None = None,
                            assets: dict | None = None) -> bytes:
    settings = normalize_settings(journal_settings)
    document = Document()
    _setup(document, settings, data)
    # A paragraph before a leading cover table anchors section/page fields in
    # LibreOffice too. Without it, PAGE fields can disappear during PDF export.
    anchor = document.add_paragraph()
    anchor.paragraph_format.space_after = Pt(0)
    anchor.paragraph_format.line_spacing = Pt(1)
    anchor.paragraph_format.keep_with_next = True
    anchor.add_run().font.size = Pt(1)
    _cover(document, data, settings, assets or {})
    if data.get('sections') or data.get('figtables') or data.get('references') or any(data.get('extra', {}).values()):
        if len(document.sections) == 1:
            _body_section(document, settings, data)
        else:
            document.add_page_break()
        _body(document, data, settings, figures)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
