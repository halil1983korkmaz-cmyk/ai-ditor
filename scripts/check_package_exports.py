"""Exercise the frozen application, including template XML and native PDFium data."""
import base64
import http.cookiejar
import io
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile

from docx import Document
from pypdf import PdfWriter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'tests'))
sys.path.insert(0, str(ROOT))
from test_formatter_templates import sample_article
from journal_templates import TEMPLATES


def main():
    executable = Path(sys.argv[1]).resolve()
    with tempfile.TemporaryDirectory(prefix='aiditor-package-check-') as folder:
        with socket.socket() as probe:
            probe.bind(('127.0.0.1', 0))
            port = probe.getsockname()[1]
        with open(Path(folder) / 'server.log', 'w+', encoding='utf-8') as log:
            proc = subprocess.Popen([str(executable), '--no-browser', '--port', str(port)],
                                    env=dict(os.environ, AIDITOR_DATA_DIR=folder), stdout=log, stderr=log)
            try:
                url = 'http://127.0.0.1:' + str(port)
                client = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
                def request(path, data=None, method=None, headers=None):
                    req = urllib.request.Request(url + path, data=data, method=method, headers=headers or {})
                    with client.open(req, timeout=30) as response:
                        return response.read()
                deadline = time.monotonic() + 60
                while time.monotonic() < deadline:
                    try:
                        if json.loads(request('/health')).get('version') == '2.2.0':
                            break
                    except (OSError, ValueError):
                        if proc.poll() is not None:
                            raise RuntimeError('Packaged process stopped before startup')
                        time.sleep(.2)
                else:
                    raise RuntimeError('Packaged service did not start')
                user = json.loads(request('/api/auth/register', json.dumps({'username': 'package_check', 'password': 'temporary-test-password', 'display_name': 'Package QA'}).encode(), 'POST', {'Content-Type': 'application/json', 'X-Aiditor-Request': '1'}))['user']
                headers = {'X-Aiditor-Account': user['id'], 'X-Aiditor-Request': '1', 'Content-Type': 'application/json'}
                pdf = PdfWriter(); pdf.add_blank_page(width=60, height=40); buffer = io.BytesIO(); pdf.write(buffer)
                setup = {'base_revision': 0, 'settings': {**next(t['settings'] for t in TEMPLATES if t['id'] == 'scholarly'), 'header_mode': 'odd_even', 'header_left': 'TEK', 'header_even_left': 'ÇİFT'}, 'assets': {'logo': {'name': 'logo.pdf', 'data': 'data:application/pdf;base64,' + base64.b64encode(buffer.getvalue()).decode()}}}
                request('/api/journal', json.dumps(setup).encode(), 'PUT', headers)
                article = sample_article(); article['cover']['ethics'] = 'Bottom-anchored package note'; article['sections'][0]['content'] = 'Örnek (2026) bu biçimi incelemiştir.'
                form = urllib.parse.urlencode({'data': json.dumps(article)}).encode()
                headers['Content-Type'] = 'application/x-www-form-urlencoded'
                result = json.loads(request('/process_docx', form, 'POST', headers))
                assert result['ok'] and result['citations']['linked_in_body'] == 1, result
                blob = request('/download_docx/' + result['key'])
                document = Document(io.BytesIO(blob))
                assert len(document.inline_shapes) == 1
                assert len(document.sections) == 2
                assert 'Bottom-anchored package note' in document.sections[0].first_page_footer._element.xml
                assert 'Bottom-anchored package note' not in document.element.xml
                assert not document.sections[1].different_first_page_header_footer
                assert request('/static/previews/scholarly.png').startswith(b'\x89PNG')
                assert 'aiditor_ref_0' in document.element.xml
                assert 'ÇİFT' in document.sections[0].even_page_header.tables[0].cell(0, 0).text
                result = json.loads(request('/process_form', form, 'POST', headers))
                with zipfile.ZipFile(io.BytesIO(request('/download/' + result['key']))) as archive:
                    assert archive.testzip() is None
                    assert r'\hyperlink{aiditor_ref_0}' in archive.read('main.tex').decode()
                print('PASS frozen package: Word header XML, PDFium logo, editable DOCX, citation target, authenticated Word/ZIP exports')
            except Exception:
                log.flush(); log.seek(0); print(log.read())
                raise
            finally:
                if os.name == 'nt' and proc.poll() is None:
                    # A one-file PyInstaller EXE has a bootloader parent and a
                    # service child. Terminate both before removing locked data.
                    subprocess.run(['taskkill', '/PID', str(proc.pid), '/T', '/F'],
                                   check=True, stdout=subprocess.DEVNULL,
                                   stderr=subprocess.DEVNULL)
                elif proc.poll() is None:
                    proc.terminate()
                try:
                    proc.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    proc.kill(); proc.wait(timeout=5)


if __name__ == '__main__':
    main()
