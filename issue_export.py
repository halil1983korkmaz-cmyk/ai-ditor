"""Journal issue (sayı): front matter pages (cover, imprint, contents) and consecutive pagination.

Front matter ("jenerik") comes from journal-level settings (imprint text, cover image, header
line) plus issue data (volume, issue, month, year) and the saved articles chosen for the issue.
Everything is produced as editable Word files.
"""
import copy
import io
import re

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

from docx_export import (ALIGN, FONTS, _custom, _element, _paragraph, _picture, _run, _table, generate_docx_from_form)
from journal_templates import normalize_settings

GRAY, LIGHT = '595959', '7F7F7F'
TOKEN = re.compile(r'\{(issn|cilt|sayi|ay|ay_en|yil)\}')


def issue_tokens(issue, settings):
    return {'issn': settings.get('issn_online') or settings.get('issn_print') or '',
            'cilt': issue.get('volume', ''), 'sayi': issue.get('issue', ''), 'yil': issue.get('year', ''),
            'ay': issue.get('month_tr', ''), 'ay_en': issue.get('month_en', '')}


def fill_line(template, issue, settings):
    values = issue_tokens(issue, settings)
    return TOKEN.sub(lambda m: str(values[m.group(1)]).strip() or '___', template)


def parse_frontmatter(text):
    """Imprint text -> list of sections.

    ``# TR / EN | 2`` starts a section (optionally 2 columns); ``#| TR / EN`` places a section beside the
    previous one; ``Name ; Institution`` are entries; ``---`` forces a new page.
    """
    sections = []
    for raw in (text or '').splitlines():
        line = raw.strip()
        if not line:
            continue
        if line == '---':
            sections.append({'break': True})
            continue
        head = re.fullmatch(r'#(\|?)\s*(.+?)(?:\s*\|\s*([12]))?', line)
        if head:
            sections.append({'title': head.group(2), 'beside': bool(head.group(1)) and bool(sections),
                             'columns': int(head.group(3) or 1), 'entries': []})
            continue
        if not sections or sections[-1].get('break'):
            sections.append({'title': '', 'beside': False, 'columns': 1, 'entries': []})
        name, _, institution = line.partition(';')
        sections[-1]['entries'].append((name.strip(), institution.strip()))
    return sections


def _base_font(doc, settings):
    name = FONTS[settings['font_family']] if settings['frontmatter_font'] == 'journal' else 'Arial'
    for style_name in ('Normal', 'Header', 'Footer'):
        style = doc.styles[style_name]
        style.font.name = name
        for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            style.element.rPr.rFonts.attrib.pop(qn('w:' + attr), None)
        style.font.size = Pt(9.5)
    doc.styles['Normal'].paragraph_format.space_after = Pt(0)
    doc.styles['Normal'].paragraph_format.space_before = Pt(0)


def _colored(p, text, size, *, color, bold=False, italic=False):
    run = _run(p, text, size, bold=bold, italic=italic, color=color)
    return run


def _line_runs(p, text, settings, size=9.5):
    """'Label: value | Label: value' with accent-coloured, bold labels like the printed masthead."""
    accent = settings['accent_color'].lstrip('#')
    for index, part in enumerate(text.split('|')):
        if index:
            _colored(p, ' | ', size, color=accent)
        label, colon, rest = part.strip().partition(':')
        if colon:
            _colored(p, label + ':', size, color=accent, bold=True)
            _colored(p, rest, size, color=GRAY)
        else:
            _colored(p, part.strip(), size, color=GRAY)


def _rule(p, edges, color):
    borders = _element('pBdr')
    for edge in edges:
        borders.append(_element(edge, val='single', sz=4, color=color, space=4))
    p._p.get_or_add_pPr().append(borders)


def _page_header(section, issue, settings, assets):
    """Masthead repeated on every imprint/contents page: logo banner and the issue line."""
    header = section.header
    header.is_linked_to_previous = False
    first = header.paragraphs[0]
    logo = assets.get('logo') if settings['show_logo'] else None
    width = round(21 - section.left_margin.cm - section.right_margin.cm, 2)
    picture_height = 0
    if logo:
        first.alignment = ALIGN['center']
        first.paragraph_format.space_after = Pt(4)
        _picture(first, logo, max_width_cm=settings['logo_width_cm'] or width, max_height_cm=12)
        picture_height = max(sum(int(n.get('cy')) for n in first._p.xpath('.//wp:extent')) / 360000, 0)
    line = header.add_paragraph()
    line.paragraph_format.space_before = Pt(2)
    line.paragraph_format.space_after = Pt(6)
    _rule(line, ('top', 'bottom'), settings['accent_color'].lstrip('#'))
    _line_runs(line, fill_line(settings['frontmatter_line'], issue, settings), settings)
    return picture_height


def _heading(container, title, settings):
    accent = settings['accent_color'].lstrip('#')
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    head, slash, tail = title.partition(' / ')
    _colored(p, head, 9.5, color=accent, bold=True)
    if slash:
        _colored(p, ' / ' + tail, 9.5, color=accent)
    return p


def _entry(container, name, institution, first=False):
    p = container.paragraphs[0] if first and container.paragraphs and not container.paragraphs[0].text else container.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0 if institution else 5)
    p.paragraph_format.keep_with_next = bool(institution)
    _colored(p, name, 9.5, color=GRAY, bold=bool(institution))
    if institution:
        q = container.add_paragraph()
        q.paragraph_format.space_after = Pt(5)
        _colored(q, institution, 8.5, color=LIGHT, italic=True)


def _section_block(container, section, settings, width):
    """Heading plus entries (1 or 2 columns) inside `container` (document or table cell)."""
    if section['title']:
        _heading(container, section['title'], settings)
    entries = section['entries']
    if section['columns'] == 2 and len(entries) > 1:
        half = (len(entries) + 1) // 2
        table = _table(container, [width / 2, width / 2])
        for cell, chunk in zip(table.rows[0].cells, (entries[:half], entries[half:])):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for index, (name, institution) in enumerate(chunk):
                _entry(cell, name, institution, first=index == 0)
    else:
        for name, institution in entries:
            _entry(container, name, institution)


def _imprint(doc, settings, width):
    sections = parse_frontmatter(settings['frontmatter_text'])
    index = 0
    while index < len(sections):
        section = sections[index]
        if section.get('break'):
            doc.add_page_break()
            index += 1
            continue
        group = [section]
        while index + len(group) < len(sections) and sections[index + len(group)].get('beside') and len(group) < 2:
            group.append(sections[index + len(group)])
        if len(group) == 1:
            _section_block(doc, section, settings, width)
        else:
            table = _table(doc, [width / 2] * 2)
            for cell, item in zip(table.rows[0].cells, group):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                _section_block(cell, {**item, 'columns': 1}, settings, width / 2)
        index += len(group)


def _anchor_full_page(run, blob):
    """Full-bleed picture behind the text, positioned against the page (Word "behind text")."""
    run.add_picture(io.BytesIO(_normalize_png(blob)), width=Cm(21), height=Cm(29.7))
    inline = run._r.xpath('.//wp:inline')[0]
    anchor = parse_xml(
        '<wp:anchor %s distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="1" '
        'locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/></wp:anchor>' % nsdecls('wp'))
    anchor.insert(3, inline.find(qn('wp:extent')))
    for tag in ('wp:docPr', 'wp:cNvGraphicFramePr', 'a:graphic'):
        anchor.append(inline.find(qn(tag)))
    inline.getparent().replace(inline, anchor)


def _normalize_png(blob):
    from PIL import Image, ImageOps
    with Image.open(io.BytesIO(blob)) as image:
        image = ImageOps.exif_transpose(image).convert('RGB')
        stream = io.BytesIO()
        image.save(stream, format='PNG')
    return stream.getvalue()


def _cover_page(doc, issue, settings, assets):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin', 'header_distance', 'footer_distance'):
        setattr(section, side, Cm(0))
    _, blob = assets['cover']
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(max(1, settings['cover_text_top_cm'] * 28.3465))
    _anchor_full_page(spacer.add_run(), blob)
    text = fill_line(settings['frontmatter_cover_line'], issue, settings)
    p = doc.add_paragraph()
    p.alignment = ALIGN['center']
    _run(p, text, settings['cover_text_size_pt'], bold=True, color=settings['cover_text_color'])


def _replace_range(texts, start, end, value):
    """Overwrite characters [start, end) of the concatenated w:t texts, keeping run formatting."""
    offset, placed = 0, False
    for node in texts:
        current = node.text or ''
        low, high = max(start - offset, 0), min(end - offset, len(current))
        if low < high or (not placed and low == high and start == offset and start == end):
            node.text = current[:low] + ('' if placed else value) + current[high:]
            node.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            placed = True
        offset += len(current)


def fill_placeholders(paragraph, values):
    """Replace runs of underscores (___, ____-____) in reading order with `values`."""
    texts = list(paragraph.iter(qn('w:t')))
    joined = ''.join(t.text or '' for t in texts)
    spans = list(re.finditer(r'_+(?:-_+)?', joined))
    for match, value in reversed(list(zip(spans, values))):
        if str(value).strip():
            _replace_range(texts, match.start(), match.end(), str(value).strip())


def _texts(paragraph):
    return list(paragraph.iter(qn('w:t')))


def _fill_toc_entry(sdt, entry):
    """The template's contents entry is one paragraph: type line, Turkish title, English title, pages."""
    paragraph = next(sdt.iter(qn('w:p')))
    nodes = _texts(paragraph)
    values = [n.text or '' for n in nodes]

    def set_text(index, value):
        nodes[index].text = value
        nodes[index].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')

    kind_tr, _, kind_en = entry['type'].partition(' / ')
    # ['___________', ' Makalesi', ' / ', '__________ ', 'Article', 'Makalenin Türkçe Adı', 'Makalenin ', 'İngilizce ', 'Adı', '___', '-', '___']
    if len(values) == 12:
        set_text(0, kind_tr.strip()); set_text(1, '')
        set_text(2, ' / ' if kind_en.strip() else ''); set_text(3, ''); set_text(4, kind_en.strip())
        set_text(5, entry['title'])
        set_text(6, entry['title_en']); set_text(7, ''); set_text(8, '')
        first, _, last = entry['pages'].partition('–')
        set_text(9, first); set_text(10, '-' if last else ''); set_text(11, last)
    else:  # customised template: fall back to the title only
        for index, node in enumerate(nodes):
            set_text(index, entry['title'] if index == 0 else '')


def _fill_template(blob, issue, settings, entries):
    doc = Document(io.BytesIO(blob))
    values = issue_tokens(issue, settings)
    cover_values = [values['issn'], values['cilt'], values['sayi'], values['yil']]
    header_values = [values['issn'], values['cilt'], values['sayi'], values['ay'], values['ay_en'], values['yil']]
    for paragraph in doc.element.body.iter(qn('w:p')):
        if 'E-ISSN' in ''.join(t.text or '' for t in _texts(paragraph)):
            fill_placeholders(paragraph, cover_values)
    for section in doc.sections:
        for part in (section.header, section.first_page_header, section.even_page_header):
            for paragraph in part._element.iter(qn('w:p')):
                if 'E-ISSN' in ''.join(t.text or '' for t in _texts(paragraph)):
                    fill_placeholders(paragraph, header_values)
    sdts = list(doc.element.body.iter(qn('w:sdt')))
    if sdts:
        anchor = sdts[0]
        content = anchor.find(qn('w:sdtContent'))
        parent = anchor.getparent()
        position = parent.index(anchor)
        for sdt in sdts:
            parent.remove(sdt)
        for entry in entries if settings['frontmatter_toc'] == 'yes' else []:
            clone = copy.deepcopy(anchor)
            _fill_toc_entry(clone, entry)
            for child in list(clone.find(qn('w:sdtContent'))):
                parent.insert(position, child)
                position += 1
    return _save(doc)


def frontmatter_docx(issue, journal_settings, assets, entries):
    """Word file with cover, imprint and contents. `entries` come from `toc_entries`.

    With an uploaded blank jenerik template (asset 'jenerik') the template itself is filled in;
    otherwise the pages are generated from the imprint text in the journal settings.
    """
    settings = normalize_settings(journal_settings)
    if assets.get('jenerik'):
        return _fill_template(assets['jenerik'][1], issue, settings, entries)
    doc = Document()
    _base_font(doc, settings)
    has_cover = settings['frontmatter_cover'] == 'yes' and bool(assets.get('cover'))
    if has_cover:
        _cover_page(doc, issue, settings, assets)
        section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    else:
        section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    left = settings['margin_left_cm'] if _custom(settings) else 2.0
    right = settings['margin_right_cm'] if _custom(settings) else 2.0
    section.left_margin, section.right_margin = Cm(left), Cm(right)
    section.header_distance = Cm(.6)
    section.footer_distance = Cm(1.2)
    section.bottom_margin = Cm(2.0)
    height = _page_header(section, issue, settings, assets)
    section.top_margin = Cm(min(12, .6 + height + 1.3))
    width = round(21 - left - right, 2)
    if has_cover:
        # The cover page must not inherit a header from the section after it.
        doc.sections[0].header.is_linked_to_previous = False
    _imprint(doc, settings, width)
    if settings['frontmatter_toc'] == 'yes' and entries:
        doc.add_page_break()
        _toc(doc, settings, entries, width)
    return _save(doc)


def _toc(doc, settings, entries, width):
    accent = settings['accent_color'].lstrip('#')
    head, slash, tail = settings['frontmatter_toc_heading'].partition(' / ')
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    _colored(p, head, 10.5, color=accent, bold=True)
    if slash:
        _colored(p, ' / ' + tail, 10.5, color=accent)
    for entry in entries:
        if entry['type']:
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            _colored(p, entry['type'], 8.5, color=LIGHT)
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.tab_stops.add_tab_stop(Cm(width), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        _colored(p, entry['title'], 10, color='262626', bold=True)
        _colored(p, '\t' + entry['pages'], 10, color='262626', bold=True)
        if entry['title_en']:
            q = doc.add_paragraph()
            q.paragraph_format.keep_with_next = bool(entry['authors'])
            _colored(q, entry['title_en'], 9, color=LIGHT, italic=True)
        last = doc.paragraphs[-1]
        if settings['toc_show_authors'] == 'yes' and entry['authors']:
            a = doc.add_paragraph()
            _colored(a, entry['authors'], 8.5, color=GRAY)
            last = a
        last.paragraph_format.space_after = Pt(9)


def _save(doc):
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def toc_entries(articles, ranges):
    """[(project data)], [(start, end)] -> display entries for the contents page."""
    result = []
    for data, (start, end) in zip(articles, ranges):
        cover = data.get('cover', {})
        title_tr, title_en = cover.get('tr_title', '').strip(), cover.get('en_title', '').strip()
        title = title_tr or title_en
        pages = f'{start}–{end}' if start and end and start != end else str(start or '')
        result.append({'type': cover.get('article_type', '').strip(), 'title': title,
                       'title_en': title_en if title_tr and title_en and title_en != title else '',
                       'authors': ', '.join(a.get('name', '').strip() for a in data.get('authors', []) if a.get('name', '').strip()),
                       'pages': pages})
    return result


def issue_article_data(data, issue, start, end):
    """Article form data with the issue's volume/issue/year and page range applied."""
    data = copy.deepcopy(data)
    cover = data.setdefault('cover', {})
    for key, source in (('volume', 'volume'), ('issue', 'issue'), ('year', 'year')):
        if str(issue.get(source, '')).strip():
            cover[key] = str(issue[source]).strip()
    cover['start_page'], cover['end_page'] = str(start), str(end)
    return data


def build_articles(items, issue, settings, assets, first_page):
    """Render every article Word file, numbering pages consecutively.

    items: [(name, form_data, figures)]. An article spans its saved page range
    (end - start + 1, at least one page); the next starts on the following page.
    Returns [{'name','docx','start','end'}].
    """
    results, cursor = [], first_page
    for name, data, figures in items:
        cover = data.get('cover', {})
        saved_start, saved_end = str(cover.get('start_page', '')).strip(), str(cover.get('end_page', '')).strip()
        span = int(saved_end) - int(saved_start) + 1 if saved_start.isdigit() and saved_end.isdigit() else 1
        start, end = cursor, cursor + max(span, 1) - 1
        docx = generate_docx_from_form(issue_article_data(data, issue, start, end), figures, settings, assets)
        results.append({'name': name, 'docx': docx, 'start': start, 'end': end})
        cursor = end + 1
    return results
